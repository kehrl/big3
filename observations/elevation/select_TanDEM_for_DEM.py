# This file takes the Microsoft Excel spreadsheets with potential TanDEM DEMs and selects 
# the best image pairs for DEM production.
#
#
# LMK, UW, 8/16/2015

import openpyxl
import os
import numpy as np
import datelib, glaclib, icefrontlib
import shapefile, shapely.geometry
import matplotlib.pyplot as plt, matplotlib

# Choose glacier (Kanger, Helheim, Jak)
glacier = 'Helheim'

# Ranges of acceptable values:
hoa_lims = [20.0,205.0]
ebase_lims = [30.0,205.0]

# Choose relative orbits
if glacier == 'Kanger':
  orbits = [163,65]
elif glacier == 'Helheim':
  orbits = [148]
elif glacier == 'Jak':
  orbits = [42,118]

############
# TanDEM-X #
############
  
# Load excel spreadsheet
excelfile=os.path.join(os.getenv("DATA_HOME"),"Elevation/TanDEM/OrderLists/"+glacier+"_TDM.xlsx")
data = openpyxl.load_workbook(excelfile)[glacier]

# Values that we want from the file
N = len(data.columns[0])-1
hoa = np.zeros(N) # height of ambiguity
ebase = np.zeros(N) # effective baseline
scenenum = np.zeros(N) # scenenumber
date = np.zeros([N,3]) # Date
relorbit = np.zeros(N) # Relative orbit 
absorbit = np.zeros(N) # Absolute orbit
order = np.zeros(N) # Are we ordering?
id = np.zeros(N)
location = []

# Load those values from the file
for column in data.columns:
  if column[0].value == 'Start Date':
    n = 0
    for cell in column[1:]:
      date[n,:] = ([int(cell.value[0:4]),int(cell.value[5:7]),int(cell.value[8:10])])  
      n = n+1
  elif column[0].value == 'Order':
    n = 0
    for cell in column[1:]:
      try:
        order[n] = int(cell.value)
      except:
        n
      n = n+1
  elif column[0].value == 'Acquisition Item ID':
    n = 0
    for cell in column[1:]:
      id[n] = (int(cell.value))
      n = n+1
  elif column[0].value == 'Relative Orbit':
    n = 0
    for cell in column[1:]:
      relorbit[n] = (int(cell.value))
      n = n+1
  elif column[0].value == 'Effective Baseline':
    n = 0
    for cell in column[1:]:
      ebase[n] = (float(cell.value))
      n = n+1
  elif column[0].value == 'Scene Number':
    n = 0
    for cell in column[1:]:
      scenenum[n] = (int(cell.value))
      n = n+1
  elif column[0].value == 'Absolute Orbit':
    n = 0
    for cell in column[1:]:
      absorbit[n] = (int(cell.value))
      n = n+1
  elif column[0].value == 'Height of Ambiguity':
    n = 0
    for cell in column[1:]:
      hoa[n] = (float(cell.value))
      n = n+1
  elif column[0].value == 'FOOTPRINT':
    for cell in column[1:]:
      location.append (str(cell.value))
del column, cell, n

# Select candidates for DEM production
cand_orbit = []
cand_ebase = []
cand_hoa = []
for i in range(0,len(date)):
  if relorbit[i] in orbits:
    cand_orbit.append(i)
  if (ebase[i] > ebase_lims[0]) and (ebase[i] < ebase_lims[1]):
    cand_ebase.append(i)
  if (abs(hoa[i]) > hoa_lims[0]) and (abs(hoa[i]) < hoa_lims[1]):    
    cand_hoa.append(i)

# Find intersection of candidate indices
cand_ind = list(set(cand_orbit).intersection(cand_ebase,cand_hoa))

# Get potential image pairs
hoa_cand = hoa[cand_ind]
ebase_cand = ebase[cand_ind]
relorbit_cand = relorbit[cand_ind]
absorbit_cand = absorbit[cand_ind]
date_cand = date[cand_ind,:]
scenenum_cand = scenenum[cand_ind]
fracyear_cand = datelib.date_to_fracyear(date[cand_ind,0],date[cand_ind,1],date[cand_ind,2])
location_cand = np.array(location)[cand_ind]
order_cand = order[cand_ind]
print len(cand_ind)

##################
# Worldview DEMs #
##################

stereopairs = shapefile.Reader(os.path.join(os.getenv("DATA_HOME"),"Elevation/Worldview/Orders/20150721_JHK_order/stereo_dg_imagery_index_CC75_N_Gr_front_vm300m"))
monopairs = shapefile.Reader(os.path.join(os.getenv("DATA_HOME"),"Elevation/Worldview/Orders/20150721_JHK_order/validpairs"))

if glacier == 'Jak':
  pt = [-185662.0,-2272488.0]
elif glacier == 'Helheim':
  pt = [305905.0,-2576918.0]
elif glacier == 'Kanger':
  pt = [493098.0,-2293282.0]
pt = shapely.geometry.Point(pt)

stereopair_dates = []
shapes = stereopairs.shapes()
recs = stereopairs.records()
for i in range(0,len(shapes)):
  polygon = shapely.geometry.asShape(shapes[i])
  if polygon.contains(pt):
    stereopair_dates.append(datelib.date_to_fracyear(float(recs[i][1][0:4]),float(recs[i][1][5:7]),float(recs[i][1][8:10])))

monopair_dates = []
shapes = monopairs.shapes()
recs = monopairs.records()
for i in range(0,len(shapes)):
  polygon = shapely.geometry.asShape(shapes[i])
  if polygon.contains(pt):
    monopair_dates.append(datelib.date_to_fracyear(float(str(recs[i][2])[0:4]),float(str(recs[i][2])[4:6]),float(str(recs[i][2])[6:8])))

plt.figure()
for day in fracyear_cand:
  plt.plot([day,day],[-4,4],'0.7')
for day in monopair_dates:
  plt.plot([day,day],[-4,4],'r',linewidth=1.5)
for day in stereopair_dates:
  plt.plot([day,day],[-4,4],'b',linewidth=1.5)  
plt.plot([fracyear_cand[0],fracyear_cand[0]],[-4,4],'0.7',label='Potential TDM')
plt.plot([monopair_dates[0],monopair_dates[0]],[-4,4],'r',linewidth=1.5,label='WV monopair')
plt.plot([stereopair_dates[0],stereopair_dates[0]],[-4,4],'b',linewidth=1.5,label='WV stereopair')
plt.plot([fracyear_cand[0],fracyear_cand[0]],[0,0],'g',linewidth=2,label='TDM Order')
plt.plot([fracyear_cand[order_cand==1],fracyear_cand[order_cand==1]],[-4,4],'g',linewidth=2)
try:
  x,y,zb,dists = glaclib.load_flowline(glacier)
  terminus_val, terminus_time = icefronts.distance_along_flowline(x,y,dists,glacier,type='icefront')
  plt.plot(terminus_time,terminus_val/1000,'k.')
except:
  print "no terminus positions"
x_formatter = matplotlib.ticker.ScalarFormatter(useOffset=False)
plt.gca().xaxis.set_major_formatter(x_formatter)
plt.title(glacier+'\n(total='+str(N)+', potential='+str(len(fracyear_cand))+', orbits='+str(orbits)+')')
plt.xlim([2011,2015.6]) 
plt.ylim([-4,4])
plt.ylabel('Terminus (km)')
plt.legend(loc=2)

#########
# Order #
#########
os.chdir("/Users/kehrl/Bigtmp/")
order_num = np.where(order == 1)[0]
fid = open(glacier+'_order_TDX.txt','w')
fid.write('Date\tAcquisitionItemID\tAbsoluteOrbit\tRelativeOrbit\tSceneNumber\tHeightOfAmbiguity\tEffectiveBaseline\n') 
for i in order_num:
  fid.write('{0}-{1:02d}-{2:02d}\t{3:3d}\t{4:3d}\t{5}\t{6}\t{7:.4f}\t{8:.4f}\n'.format(int(date[i,0]),(int(date[i,1])),(int(date[i,2])),int(id[i]),int(absorbit[i]),int(relorbit[i]),int(scenenum[i]),hoa[i],ebase[i]))
fid.close() 